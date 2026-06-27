"""
Generate SafeRoute Bengaluru PPT Content as a DOCX file.
Each section = one PPT slide. Content is concise and presentation-ready.
"""

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import os

doc = Document()

# ── Global styles ──
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(12)
font.color.rgb = RGBColor(0x33, 0x33, 0x33)

# Helper: add a slide title
def add_slide_title(doc, slide_num, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(f"SLIDE {slide_num}")
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x00, 0x4E, 0x3E)
    
    heading = doc.add_heading(title, level=1)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0x00, 0x4E, 0x3E)
        run.font.size = Pt(24)

def add_subheading(doc, text):
    heading = doc.add_heading(text, level=2)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
        run.font.size = Pt(16)

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.bold = True
        run.font.size = Pt(12)
        p.add_run(text)
    else:
        p.add_run(text)

def add_body(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)

def add_page_break(doc):
    doc.add_page_break()


# ════════════════════════════════════════════════════════════════
# SLIDE 1 — PROBLEM STATEMENT
# ════════════════════════════════════════════════════════════════
add_slide_title(doc, 1, "Problem Statement")

add_body(doc,
    "Urban women in India face a critical safety gap during daily commutes. "
    "According to the NCRB (National Crime Records Bureau) 2023 report, India recorded over 4.45 lakh cases of crimes against women — "
    "a 4% year-on-year rise. Bengaluru alone reported 2,000+ street harassment and assault cases in 2023. "
    "Yet, the navigation tools used by millions — Google Maps, Ola Maps, Apple Maps — optimize routes "
    "exclusively for time and distance, completely ignoring personal safety."
)

add_subheading(doc, "The Core Problem")

add_bullet(doc, "Existing navigation apps treat all roads equally — a well-lit, CCTV-monitored MG Road is scored the same as an isolated, unlit back lane in Majestic at midnight.")
add_bullet(doc, "Women, night-shift workers, college students, and elderly citizens have no data-driven way to choose safer routes.")
add_bullet(doc, "There is no real-time mechanism to alert users when they enter a danger zone or to crowdsource live safety intelligence from the community.")
add_bullet(doc, "SOS/emergency features are either absent or disconnected from the navigation workflow in existing apps.")

add_subheading(doc, "Who Is Affected?")

add_bullet(doc, "Women IT/BPO employees", bold_prefix="")
add_bullet(doc, " — working night shifts in Whitefield, Electronic City, and Manyata Tech Park.")
add_bullet(doc, "College students", bold_prefix="")
add_bullet(doc, " — traveling late from libraries, labs, and campus events.")
add_bullet(doc, "Tourists", bold_prefix="")
add_bullet(doc, " — unfamiliar with city geography and high-risk zones.")
add_bullet(doc, "Elderly citizens", bold_prefix="")
add_bullet(doc, " — needing proximity to medical and police infrastructure.")

add_subheading(doc, "Key Statistic")
add_body(doc, "Over 66% of women in Indian metros report feeling unsafe during night-time commutes (Thomson Reuters Foundation Survey, 2023).")

add_page_break(doc)


# ════════════════════════════════════════════════════════════════
# SLIDE 2 — EXISTING SOLUTIONS AND GAPS
# ════════════════════════════════════════════════════════════════
add_slide_title(doc, 2, "Existing Solutions & Gaps")

add_subheading(doc, "Current Navigation Landscape")
add_body(doc, "Today's navigation ecosystem is dominated by a few major players. While they excel at time-and-distance optimization, none address personal safety as a routing parameter.")

# Table: Feature comparison
table = doc.add_table(rows=8, cols=4)
table.style = 'Medium Shading 1 Accent 1'
table.alignment = WD_TABLE_ALIGNMENT.CENTER

headers = ['Feature', 'Google Maps', 'Ola Maps', 'SafeRoute']
for i, h in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.text = h
    for p in cell.paragraphs:
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(10)

rows_data = [
    ['Shortest / Fastest Route',    '✅', '✅', '✅'],
    ['Safety-First Routing',        '❌', '❌', '✅'],
    ['CCTV-Aware Routing',          '❌', '❌', '✅'],
    ['Crime Zone Avoidance',        '❌', '❌', '✅'],
    ['Real-Time Danger Alerts',     '❌', '❌', '✅'],
    ['SOS with Live GPS to Contacts','❌', '❌', '✅'],
    ['Community Safety Feed',       '❌', '❌', '✅'],
]
for r_idx, row_data in enumerate(rows_data, start=1):
    for c_idx, val in enumerate(row_data):
        table.rows[r_idx].cells[c_idx].text = val
        for p in table.rows[r_idx].cells[c_idx].paragraphs:
            for r in p.runs:
                r.font.size = Pt(10)

add_body(doc, "")  # spacer

add_subheading(doc, "Gaps in Existing Solutions")
add_bullet(doc, "No safety parameterization: ", bold_prefix="")
add_body(doc, "Routes are computed purely on travel time/distance. No integration of CCTV density, streetlight data, crowd activity, or crime history.")
add_bullet(doc, "No real-time safety alerts: ", bold_prefix="")
add_body(doc, "Users receive no warning when entering a low-safety zone or a crime hotspot.")
add_bullet(doc, "No community intelligence: ", bold_prefix="")
add_body(doc, "No mechanism for citizens to report broken streetlights, suspicious activity, or road hazards that feed back into routing decisions.")
add_bullet(doc, "Disconnected emergency response: ", bold_prefix="")
add_body(doc, "SOS features (if any) are not integrated with navigation — they exist as standalone apps (e.g., Nirbhaya app) with poor adoption.")

add_page_break(doc)


# ════════════════════════════════════════════════════════════════
# SLIDE 3 — PROPOSED SOLUTION
# ════════════════════════════════════════════════════════════════
add_slide_title(doc, 3, "Proposed Solution — SafeRoute Bengaluru")

add_body(doc,
    "SafeRoute Bengaluru is an AI-powered, safety-first navigation platform that computes the safest practical route — "
    "not just the shortest or fastest — by analyzing real-time and historical safety data across Bengaluru's road network."
)

add_subheading(doc, "How It Works")
add_bullet(doc, "Every road segment in Bengaluru is assigned a dynamic Safety Score (0–100) computed from 5 weighted parameters: Streetlight density (35%), CCTV coverage (25%), Crowd/pedestrian activity (20%), Emergency facility proximity (20%), and Crime history penalty.")
add_bullet(doc, "A modified Dijkstra/Yen's K-Shortest Paths algorithm uses this Safety Score as the primary edge weight instead of just distance, producing 3 route options: Safest, Balanced, and Fastest.")
add_bullet(doc, "Time-of-day multipliers automatically adjust scores — routes at 2 AM are penalized more heavily on poorly-lit roads than the same routes at 2 PM.")

add_subheading(doc, "Core Modules Built")
add_bullet(doc, "Safety-Weighted Routing Engine — Graph-based routing using NetworkX with Yen's K-Shortest Paths, SciPy cKDTree for GPS snap, and GeoJSON polyline rendering.", bold_prefix="1. ")
add_bullet(doc, "Interactive Map with Live Navigation — Leaflet.js map with turn-by-turn GPS navigation HUD, real-time progress tracking, and route simulation.", bold_prefix="2. ")
add_bullet(doc, "SOS Emergency System — One-tap SOS with 3-second abort countdown, auto-sends SMS with live GPS coordinates to emergency contacts via Twilio API.", bold_prefix="3. ")
add_bullet(doc, "Community Safety Feed — Crowdsourced geo-tagged safety reports (harassment, broken lights, suspicious activity) with verification system.", bold_prefix="4. ")
add_bullet(doc, "Real-Time Danger Alerts — WebSocket-based push notifications when user enters road segments with safety score below 50.", bold_prefix="5. ")
add_bullet(doc, "Live Safety Heatmap — Color-coded overlay showing safe (green) and unsafe (red) zones across the city.", bold_prefix="6. ")

add_page_break(doc)


# ════════════════════════════════════════════════════════════════
# SLIDE 4 — UNIQUE SELLING PROPOSITION (USP)
# ════════════════════════════════════════════════════════════════
add_slide_title(doc, 4, "Unique Selling Proposition (USP)")

add_body(doc,
    "SafeRoute is not just another navigation app — it is the first platform to treat personal safety as a first-class routing parameter backed by real data."
)

add_subheading(doc, "What Makes SafeRoute Unique")

add_bullet(doc, "Multi-Parameter Safety Score — ", bold_prefix="🛡️ ")
add_body(doc, "Unlike any existing nav app, SafeRoute computes a composite Safety Score per road segment using 5 real-world data dimensions: CCTV density, streetlight coverage, crowd activity, emergency facility proximity, and crime history. This is not a binary safe/unsafe label — it is a granular 0–100 score with full breakdown transparency.")

add_bullet(doc, "Safety-Cost Routing Algorithm — ", bold_prefix="🧮 ")
add_body(doc, "Our modified Dijkstra cost function: C = L × (1 + β × (1 − S/100)). When β=10 (Safest mode), the router adds up to a 10× virtual length penalty to detour around unsafe roads. No other navigation app uses safety as a graph edge weight.")

add_bullet(doc, "Time-Aware Dynamic Scoring — ", bold_prefix="🌙 ")
add_body(doc, "Safety scores automatically adjust based on time-of-day. The same road at 3 PM (multiplier 1.0) scores very differently at 1 AM (multiplier 0.65). Night-shift workers get routes optimized for late-night safety.")

add_bullet(doc, "Integrated Emergency Response — ", bold_prefix="🚨 ")
add_body(doc, "SOS is not a separate app — it is embedded inside navigation. One tap triggers a 3-second countdown, then auto-dispatches SMS with live GPS to all emergency contacts via Twilio. No other routing app has an integrated SOS pipeline.")

add_bullet(doc, "Community-Powered Intelligence — ", bold_prefix="📢 ")
add_body(doc, "Users can report live hazards (harassment, broken streetlights, suspicious groups) that are geo-tagged and verified by the community, feeding directly back into the Safety Score engine in real time.")

add_bullet(doc, "Built for Bengaluru — ", bold_prefix="🏙️ ")
add_body(doc, "Hyper-local: Uses Bengaluru's OpenStreetMap road network, NCRB Karnataka crime data, Bengaluru Safe City CCTV data (5.3 lakh cameras), BBMP streetlight records, and BMTC/BMRCL transit stops. Not a generic global product.")

add_page_break(doc)


# ════════════════════════════════════════════════════════════════
# SLIDE 5 — SYSTEM ARCHITECTURE
# ════════════════════════════════════════════════════════════════
add_slide_title(doc, 5, "System Architecture")

add_body(doc,
    "SafeRoute follows a 3-tier architecture: Client Layer → API Gateway → Data Layer, "
    "with real-time capabilities via WebSockets and background processing via Celery workers."
)

add_subheading(doc, "Architecture Overview")
add_body(doc, """
┌─────────────────────────────────────────────────────────────┐
│                    CLIENT LAYER                              │
│    Next.js 16 (React 19) + Leaflet.js + Tailwind CSS v4    │
│    [Map View] [Route Cards] [SOS Button] [Safety Feed]      │
└────────────────────────┬────────────────────────────────────┘
                         │  REST API (JSON) + WebSockets
┌────────────────────────▼────────────────────────────────────┐
│                   API GATEWAY (FastAPI)                       │
│  ┌──────────┐ ┌──────────┐ ┌────────┐ ┌──────────────────┐  │
│  │Auth (JWT │ │ Routing  │ │  SOS   │ │ Report Ingestion │  │
│  │  + OTP)  │ │ Engine   │ │Service │ │    Service       │  │
│  └──────────┘ └──────────┘ └────────┘ └──────────────────┘  │
│  ┌──────────────────┐  ┌────────────────────────────────┐   │
│  │Safety Score Engine│  │ WebSocket Manager (Live Alerts)│   │
│  └──────────────────┘  └────────────────────────────────┘   │
└────────────────────────┬────────────────────────────────────┘
                         │
┌────────────────────────▼────────────────────────────────────┐
│                     DATA LAYER                               │
│  ┌──────────────────┐  ┌────────┐  ┌──────────────────────┐ │
│  │PostgreSQL+PostGIS│  │ Redis  │  │ Celery Workers       │ │
│  │(Road Graph,Scores│  │(Cache, │  │(Score Refresh,       │ │
│  │ Crime, CCTV)     │  │Broker) │  │ Graph Rebuild)       │ │
│  └──────────────────┘  └────────┘  └──────────────────────┘ │
└─────────────────────────────────────────────────────────────┘
""")

add_subheading(doc, "Data Flow")
add_bullet(doc, "User enters source & destination on the Next.js frontend.")
add_bullet(doc, "Frontend sends POST request to /api/v1/route/options on FastAPI backend.")
add_bullet(doc, "Backend snaps GPS coordinates to the nearest graph node using SciPy cKDTree.")
add_bullet(doc, "Yen's K-Shortest Paths algorithm computes 3 distinct routes ranked by safety-weighted cost.")
add_bullet(doc, "Each route's GeoJSON polyline is returned with per-segment safety scores for color-coded rendering.")
add_bullet(doc, "WebSocket connection streams real-time danger alerts when user GPS hits a low-safety segment (score < 50).")
add_bullet(doc, "SOS trigger sends SMS via Twilio to stored emergency contacts with live Google Maps GPS link.")

add_subheading(doc, "Containerized Infrastructure")
add_body(doc, "Docker Compose orchestrates 4 services: PostgreSQL/PostGIS database, Redis cache/broker, FastAPI application server, and Celery worker/beat for background tasks (nightly score refresh, weekly graph rebuild).")

add_page_break(doc)


# ════════════════════════════════════════════════════════════════
# SLIDE 6 — TECH STACK
# ════════════════════════════════════════════════════════════════
add_slide_title(doc, 6, "Technology Stack")

# Table: Tech stack
table2 = doc.add_table(rows=13, cols=3)
table2.style = 'Medium Shading 1 Accent 1'
table2.alignment = WD_TABLE_ALIGNMENT.CENTER

headers2 = ['Layer', 'Technology', 'Purpose']
for i, h in enumerate(headers2):
    cell = table2.rows[0].cells[i]
    cell.text = h
    for p in cell.paragraphs:
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(10)

stack_data = [
    ['Frontend Framework', 'Next.js 16 (React 19)', 'Server-side rendering, App Router, fast interactive UI'],
    ['Styling', 'Tailwind CSS v4', 'Utility-first responsive design system'],
    ['Map Engine', 'Leaflet.js + react-leaflet', 'Interactive map rendering, polyline overlays, GPS tracking'],
    ['Animation', 'Framer Motion', 'Smooth UI transitions and micro-animations'],
    ['Backend Framework', 'FastAPI (Python 3.11)', 'High-performance async API server with auto-generated docs'],
    ['Database', 'PostgreSQL 15 + PostGIS 3.3', 'Geospatial queries (ST_DWithin), road graph storage'],
    ['Cache / Broker', 'Redis 7', 'Safety score caching (5-min TTL), Celery task broker'],
    ['Graph Engine', 'NetworkX + SciPy cKDTree', 'Directed graph routing, O(log n) nearest-node lookup'],
    ['Background Jobs', 'Celery + Celery Beat', 'Async score refresh, graph rebuild scheduling'],
    ['ORM / Migrations', 'SQLAlchemy 2.0 + GeoAlchemy2 + Alembic', 'Async ORM, spatial columns, versioned migrations'],
    ['SMS / Emergency', 'Twilio SMS API', 'SOS alert dispatch to emergency contacts'],
    ['Infrastructure', 'Docker Compose (4 containers)', 'Reproducible local dev: DB, Redis, API, Celery worker'],
]
for r_idx, row_data in enumerate(stack_data, start=1):
    for c_idx, val in enumerate(row_data):
        table2.rows[r_idx].cells[c_idx].text = val
        for p in table2.rows[r_idx].cells[c_idx].paragraphs:
            for r in p.runs:
                r.font.size = Pt(9)

add_body(doc, "")
add_subheading(doc, "Key Libraries & APIs")
add_bullet(doc, "OSMnx — Downloads and models OpenStreetMap road networks for Bengaluru.", bold_prefix="")
add_bullet(doc, "Shapely — Geometric coordinate calculations and spatial manipulation.", bold_prefix="")
add_bullet(doc, "GeoJSON — Standard format for route polylines and heatmap overlays.", bold_prefix="")
add_bullet(doc, "python-jose + bcrypt — JWT authentication with OTP-based phone verification.", bold_prefix="")
add_bullet(doc, "Zustand — Lightweight React state management for route selections.", bold_prefix="")
add_bullet(doc, "HTML5 Geolocation API — Real-time GPS tracking from user's browser.", bold_prefix="")

add_page_break(doc)


# ════════════════════════════════════════════════════════════════
# SLIDE 7 — IMPACT AND BENEFITS
# ════════════════════════════════════════════════════════════════
add_slide_title(doc, 7, "Impact & Benefits")

add_subheading(doc, "Social Impact")
add_bullet(doc, "Directly addresses women's safety — one of India's most pressing urban challenges — through technology-driven, data-backed route optimization.", bold_prefix="")
add_bullet(doc, "Empowers vulnerable populations (women, elderly, students, tourists) with actionable safety intelligence before and during their commute.", bold_prefix="")
add_bullet(doc, "Creates a participatory safety network where citizens actively contribute to making their city safer through community reports.", bold_prefix="")

add_subheading(doc, "Measurable Benefits")

table3 = doc.add_table(rows=6, cols=2)
table3.style = 'Medium Shading 1 Accent 1'
table3.alignment = WD_TABLE_ALIGNMENT.CENTER

headers3 = ['Metric', 'Target (3 months post-launch)']
for i, h in enumerate(headers3):
    cell = table3.rows[0].cells[i]
    cell.text = h
    for p in cell.paragraphs:
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(10)

metrics = [
    ['Routes computed per day', '500+'],
    ['Avg safety score of recommended routes', '> 72 / 100'],
    ['SOS activations resolved within 5 min', '> 90%'],
    ['User-reported incidents resolved', '> 80%'],
    ['DAU / MAU ratio', '> 0.25'],
]
for r_idx, row_data in enumerate(metrics, start=1):
    for c_idx, val in enumerate(row_data):
        table3.rows[r_idx].cells[c_idx].text = val
        for p in table3.rows[r_idx].cells[c_idx].paragraphs:
            for r in p.runs:
                r.font.size = Pt(10)

add_body(doc, "")
add_subheading(doc, "Stakeholder Benefits")
add_bullet(doc, "For Bengaluru Police: Real-time crowd-sourced incident data and hotspot analytics can guide patrol deployment.", bold_prefix="")
add_bullet(doc, "For BBMP/Smart City: Safety score data identifies where to prioritize streetlight installation and CCTV expansion.", bold_prefix="")
add_bullet(doc, "For Employers: Organizations with night-shift workforces (IT, BPO, hospitals) can integrate SafeRoute to ensure employee commute safety.", bold_prefix="")
add_bullet(doc, "For Citizens: A transparent, explainable safety scoring system — users see exactly why a route is rated safe or unsafe (lighting score, CCTV count, crime history).", bold_prefix="")

add_page_break(doc)


# ════════════════════════════════════════════════════════════════
# SLIDE 8 — FUTURE SCOPE
# ════════════════════════════════════════════════════════════════
add_slide_title(doc, 8, "Future Scope")

add_subheading(doc, "Phase 2 — Advanced Routing (Next 3 Months)")
add_bullet(doc, "A* Search with Safety Heuristic — Replace Dijkstra with A* using a composite heuristic that considers both Euclidean distance and regional safety, achieving 3–5× faster route computation.", bold_prefix="")
add_bullet(doc, "Multi-Objective Optimization (NSGA-II) — Return Pareto-optimal routes that simultaneously optimize safety, distance, and time, with a user-controlled Safety ↔ Speed slider.", bold_prefix="")
add_bullet(doc, "Night Mode — Automatic activation after 9 PM with stricter safety thresholds, avoiding isolated roads even if shorter.", bold_prefix="")

add_subheading(doc, "Phase 3 — ML & Computer Vision (6 Months)")
add_bullet(doc, "Satellite Night-Light Analysis — Use NASA VIIRS night-time light data to automatically compute lighting scores for every road segment without manual surveys.", bold_prefix="")
add_bullet(doc, "Street-Level Image Classification — ML model trained on Google Street View / Mapillary images to classify roads as well-lit, moderately-lit, or dark.", bold_prefix="")
add_bullet(doc, "Predictive Crime Scoring — Time-series forecasting of crime hotspots using NCRB + news scraping data with LSTM/Transformer models.", bold_prefix="")

add_subheading(doc, "Phase 4 — Platform Expansion (12 Months)")
add_bullet(doc, "Flutter Mobile App — Native Android/iOS app with background GPS tracking, push notifications, and offline route caching.", bold_prefix="")
add_bullet(doc, "Multi-City Expansion — Scale to other Indian metros: Delhi, Mumbai, Hyderabad, Chennai — using the same data pipeline with city-specific CCTV and crime data.", bold_prefix="")
add_bullet(doc, "Public API & Smart City Integration — Expose safety score API for integration with Ola, Uber, Namma Yatri, and Smart City dashboards.", bold_prefix="")
add_bullet(doc, "Wearable Integration — Connect with smartwatches for silent SOS trigger (double-press power button) and haptic danger alerts.", bold_prefix="")

add_subheading(doc, "Phase 5 — Ecosystem (18+ Months)")
add_bullet(doc, "Corporate Safety Dashboard — B2B product for IT parks to monitor employee commute safety in real time.", bold_prefix="")
add_bullet(doc, "Government MoU Integration — Direct data exchange with Bengaluru Safe City, Karnataka Police, BBMP for real-time CCTV and incident feeds.", bold_prefix="")
add_bullet(doc, "Gamification & Rewards — Incentivize community reporting with safety points, badges, and partnerships with local businesses.", bold_prefix="")


# ── Save ──
output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "SafeRoute_PPT_Content.docx")
doc.save(output_path)
print(f"DOCX saved to: {output_path}")
